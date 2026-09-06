# -*- coding: utf-8 -*-
"""
生成 Excel 分析报表（excel/电商销售分析报告.xlsx）。

内容：
  概览        —— KPI 指标卡
  月度趋势    —— 折线图
  品类分析    —— 条形图
  地域分析    —— 条形图
  RFM分层     —— 条形图
  渠道分析    —— 条形图
  产品TOP10   —— 条形图

依赖：openpyxl
运行方式：python scripts/excel_report.py
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "data" / "processed"
EXCEL_DIR = BASE / "excel"

FONT = "微软雅黑"
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name=FONT, size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color="2F5597")
KPI_FONT = Font(name=FONT, bold=True, size=14, color="2F5597")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_dataframe(ws, df, start_row=2, money_cols=(), pct_cols=(), int_cols=()):
    """把 DataFrame 写入工作表，返回数据区末行行号。"""
    # 表头
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=str(col))
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    # 数据
    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, col in enumerate(df.columns, start=1):
            v = row[col]
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = Alignment(horizontal="center" if j == 1 else "right")
            if col in money_cols:
                c.number_format = "#,##0.00"
            elif col in pct_cols:
                c.number_format = "0.00"
            elif col in int_cols:
                c.number_format = "#,##0"
    # 列宽
    for j, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(x)) for x in df[col].head(20)]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width * 1.6, 10), 28)
    return start_row + len(df)


def add_bar_chart(ws, data_df, categories_col, values_col, title, anchor):
    chart = BarChart()
    chart.type = "bar"  # 横向条形图，适合品类名较长
    chart.style = 10
    chart.title = title
    chart.height = 10
    chart.width = 18
    last_row = ws.max_row
    data = Reference(ws, min_col=ws.max_column, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    ws.add_chart(chart, anchor)
    return chart


def add_line_chart(ws, data_df, cats_col, values_col, title, anchor):
    chart = LineChart()
    chart.title = title
    chart.style = 10
    chart.height = 10
    chart.width = 20
    last_row = ws.max_row
    data = Reference(ws, min_col=2, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    ws.add_chart(chart, anchor)
    return chart


def build_overview(ws, kpi):
    """KPI 指标卡。"""
    ws.sheet_view.showGridLines = False
    ws["B2"] = "电商销售数据分析报告"
    ws["B2"].font = TITLE_FONT
    ws["B3"] = "数据周期：2024-01-01 ~ 2025-12-31  ·  数据来源：模拟订单数据"
    ws["B3"].font = Font(name=FONT, size=10, color="808080")

    # 指标卡横向排布
    start_col = 2
    for _, row in kpi.iterrows():
        metric, value = row["metric"], row["value"]
        c = ws.cell(row=6, column=start_col, value=metric)
        c.font = Font(name=FONT, size=10, color="808080")
        c.alignment = Alignment(horizontal="center")
        v = ws.cell(row=7, column=start_col, value=value)
        v.font = KPI_FONT
        v.alignment = Alignment(horizontal="center")
        v.number_format = "#,##0.00" if isinstance(value, float) else "#,##0"
        start_col += 3

    # 边框美化
    for col in range(2, start_col, 3):
        for r in (6, 7):
            ws.cell(row=r, column=col).border = Border(bottom=Side(style="thin", color="2F5597"))
    ws.column_dimensions["A"].width = 2


def main():
    EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    monthly = pd.read_csv(OUT / "monthly_trend.csv", encoding="utf-8-sig")
    cat = pd.read_csv(OUT / "category_summary.csv", encoding="utf-8-sig")
    reg = pd.read_csv(OUT / "region_summary.csv", encoding="utf-8-sig")
    rfm = pd.read_csv(OUT / "rfm_segments.csv", encoding="utf-8-sig")
    ch = pd.read_csv(OUT / "channel_summary.csv", encoding="utf-8-sig")
    top = pd.read_csv(OUT / "product_top10.csv", encoding="utf-8-sig")
    kpi = pd.read_csv(OUT / "kpi_summary.csv", encoding="utf-8-sig")

    wb = Workbook()

    # 概览
    ws = wb.active
    ws.title = "概览"
    build_overview(ws, kpi)

    # 月度趋势
    ws = wb.create_sheet("月度趋势")
    write_dataframe(ws, monthly, money_cols=("gmv", "avg_order_value"),
                    pct_cols=("mom_pct", "yoy_pct"), int_cols=("order_cnt",))
    add_line_chart(ws, monthly, "ym", "gmv", "月度 GMV 趋势", "A30")

    # 品类分析
    ws = wb.create_sheet("品类分析")
    write_dataframe(ws, cat, money_cols=("revenue", "cost", "gross_profit"),
                    pct_cols=("margin_pct", "revenue_share_pct"), int_cols=("order_cnt", "sold_qty"))
    add_bar_chart(ws, cat, "category", "revenue", "各品类销售额", "A12")

    # 地域分析
    ws = wb.create_sheet("地域分析")
    write_dataframe(ws, reg, money_cols=("revenue",), int_cols=("order_cnt", "buyer_cnt"))
    add_bar_chart(ws, reg, "province", "revenue", "各省份销售额 TOP", "A20")

    # RFM 分层
    ws = wb.create_sheet("RFM分层")
    write_dataframe(ws, rfm, money_cols=("avg_monetary",), pct_cols=("pct",), int_cols=("cnt",))
    add_bar_chart(ws, rfm, "segment", "cnt", "客户分层人数", "A10")

    # 渠道分析
    ws = wb.create_sheet("渠道分析")
    write_dataframe(ws, ch, money_cols=("revenue", "avg_order_value"), int_cols=("order_cnt",))
    add_bar_chart(ws, ch, "channel", "revenue", "各渠道销售额", "A10")

    # 产品 TOP10
    ws = wb.create_sheet("产品TOP10")
    write_dataframe(ws, top, money_cols=("revenue", "cost"),
                    pct_cols=("margin_pct",), int_cols=("sold_qty",))
    add_bar_chart(ws, top, "product_name", "revenue", "畅销产品 TOP10", "A15")

    out_path = EXCEL_DIR / "电商销售分析报告.xlsx"
    wb.save(out_path)
    print(f"已生成：{out_path}")


if __name__ == "__main__":
    main()
